"""
Excel Report Generator for IT Support Dashboard
Author: Yash Somaiya
Generates a formatted Excel workbook from ticket data — mirrors real ServiceNow/Excel reporting workflows
"""

import csv
import json
import os
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.series import DataPoint
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    print("[WARN] openpyxl not installed. Run: pip install openpyxl")
    print("[INFO] Run ticketing_dashboard.py first to generate CSV data, then install openpyxl and re-run this script.")

def load_tickets(csv_path="data/tickets_export.csv"):
    tickets = []
    with open(csv_path, newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            tickets.append(row)
    return tickets

def load_summary(json_path="reports/dashboard_summary.json"):
    with open(json_path) as f:
        return json.load(f)

def style_header_cell(cell, bg_color="1F4E79", font_color="FFFFFF"):
    cell.font = Font(bold=True, color=font_color, size=11)
    cell.fill = PatternFill("solid", fgColor=bg_color)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def build_excel_report(tickets, stats, output_path="reports/IT_Support_Dashboard.xlsx"):
    if not OPENPYXL_AVAILABLE:
        print("[SKIP] Excel generation skipped — openpyxl not available.")
        return

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary Dashboard ──────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary Dashboard"
    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 18
    ws_sum.column_dimensions["C"].width = 28
    ws_sum.column_dimensions["D"].width = 18

    # Title
    ws_sum.merge_cells("A1:D1")
    title_cell = ws_sum["A1"]
    title_cell.value = "IT Support Ticketing Dashboard"
    title_cell.font = Font(bold=True, size=16, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center")
    ws_sum.row_dimensions[1].height = 30

    ws_sum.merge_cells("A2:D2")
    sub_cell = ws_sum["A2"]
    sub_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Author: Yash Somaiya"
    sub_cell.font = Font(italic=True, size=10, color="666666")
    sub_cell.alignment = Alignment(horizontal="center")

    # KPI section
    kpis = [
        ("Total Tickets",       stats["total_tickets"],          "Open Tickets",         stats["open_tickets"]),
        ("Resolved Tickets",    stats["resolved_tickets"],        "Resolution Rate",      f"{stats['resolution_rate_pct']}%"),
        ("SLA Compliance",      f"{stats['sla_compliance_pct']}%","Avg Resolution Time", f"{stats['avg_resolution_hrs']} hrs"),
    ]
    ws_sum.append([])
    for row_data in kpis:
        row = ws_sum.max_row + 1
        ws_sum.cell(row=row, column=1, value=row_data[0]).font = Font(bold=True, color="1F4E79")
        ws_sum.cell(row=row, column=2, value=row_data[1]).font = Font(bold=True, size=13)
        ws_sum.cell(row=row, column=3, value=row_data[2]).font = Font(bold=True, color="1F4E79")
        ws_sum.cell(row=row, column=4, value=row_data[3]).font = Font(bold=True, size=13)
        ws_sum.row_dimensions[row].height = 22

    # Tickets by Category table
    ws_sum.append([])
    ws_sum.append([])
    header_row = ws_sum.max_row
    ws_sum.cell(row=header_row, column=1, value="Category")
    ws_sum.cell(row=header_row, column=2, value="Ticket Count")
    style_header_cell(ws_sum.cell(row=header_row, column=1))
    style_header_cell(ws_sum.cell(row=header_row, column=2))

    cat_start = header_row + 1
    for cat, count in stats["by_category"].items():
        r = ws_sum.max_row + 1
        ws_sum.cell(row=r, column=1, value=cat).border = thin_border()
        count_cell = ws_sum.cell(row=r, column=2, value=count)
        count_cell.alignment = Alignment(horizontal="center")
        count_cell.border = thin_border()
    cat_end = ws_sum.max_row

    # Bar chart for categories
    chart = BarChart()
    chart.type = "col"
    chart.title = "Tickets by Category"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Category"
    chart.style = 10
    chart.width = 18
    chart.height = 10

    data_ref   = Reference(ws_sum, min_col=2, min_row=header_row, max_row=cat_end)
    cats_ref   = Reference(ws_sum, min_col=1, min_row=cat_start, max_row=cat_end)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws_sum.add_chart(chart, "F4")

    # ── Sheet 2: All Tickets ────────────────────────────────────────────────────
    ws_tickets = wb.create_sheet("All Tickets")
    headers = list(tickets[0].keys())
    col_widths = [12, 18, 18, 16, 16, 14, 38, 14, 18, 18, 10, 40, 12]

    for i, h in enumerate(headers, 1):
        cell = ws_tickets.cell(row=1, column=i, value=h.replace("_", " ").title())
        style_header_cell(cell)
        ws_tickets.column_dimensions[get_column_letter(i)].width = col_widths[i-1] if i <= len(col_widths) else 14

    for row_idx, ticket in enumerate(tickets, 2):
        for col_idx, key in enumerate(headers, 1):
            cell = ws_tickets.cell(row=row_idx, column=col_idx, value=ticket[key])
            cell.border = thin_border()
            cell.alignment = Alignment(vertical="center")

            # Colour code by priority
            if key == "priority":
                colors = {
                    "P1 - Critical": "FF4444",
                    "P2 - High":     "FF9933",
                    "P3 - Medium":   "FFCC00",
                    "P4 - Low":      "99CC66",
                }
                color = colors.get(ticket[key], "FFFFFF")
                cell.fill = PatternFill("solid", fgColor=color)

            # Colour code SLA
            if key == "met_sla":
                if ticket[key] == "Yes":
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
                elif ticket[key] == "No":
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")

    ws_tickets.freeze_panes = "A2"
    ws_tickets.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── Sheet 3: Agent Performance ──────────────────────────────────────────────
    ws_agent = wb.create_sheet("Agent Performance")
    ws_agent.column_dimensions["A"].width = 22
    ws_agent.column_dimensions["B"].width = 18
    ws_agent.column_dimensions["C"].width = 22

    ws_agent.cell(row=1, column=1, value="Agent Performance Summary").font = Font(bold=True, size=13, color="1F4E79")
    ws_agent.append([])

    a_headers = ["Agent Name", "Tickets Handled", "% of Total"]
    for i, h in enumerate(a_headers, 1):
        style_header_cell(ws_agent.cell(row=3, column=i))
        ws_agent.cell(row=3, column=i).value = h

    total = stats["total_tickets"]
    for agent, count in stats["by_agent"].items():
        r = ws_agent.max_row + 1
        ws_agent.cell(row=r, column=1, value=agent).border = thin_border()
        ws_agent.cell(row=r, column=2, value=count).border = thin_border()
        pct = ws_agent.cell(row=r, column=3, value=f"{round(count/total*100,1)}%")
        pct.border = thin_border()
        pct.alignment = Alignment(horizontal="center")

    wb.save(output_path)
    print(f"  [OK] Excel report saved: {output_path}")
    print(f"       Sheets: Summary Dashboard | All Tickets | Agent Performance")

def main():
    print("\n" + "=" * 60)
    print("  Excel Report Generator — IT Support Dashboard")
    print("=" * 60 + "\n")

    if not os.path.exists("data/tickets_export.csv"):
        print("[ERROR] Run ticketing_dashboard.py first to generate ticket data.")
        return

    print("Loading ticket data and summary...")
    tickets = load_tickets()
    stats   = load_summary()

    print("Building Excel workbook...")
    build_excel_report(tickets, stats)

    print("\nDone. Open reports/IT_Support_Dashboard.xlsx to view.")

if __name__ == "__main__":
    main()
